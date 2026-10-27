"""
readers.py — Lectores de TXT / DOCX / EPUB y detección de capítulos
             (por estilos de Word, marcadores de texto, Calibre o spine EPUB).
"""
import codecs
import os
import re
import shutil
import subprocess
import tempfile
from pathlib import Path

from docx import Document
from ebooklib import epub
from bs4 import BeautifulSoup

from . import config as C
from .text import clean

# ==========================================================
# TXT
# ==========================================================

def _read_txt_raw(path: str) -> str:
    """Lee TXT respetando BOM UTF-8/UTF-16 y sin romper acentos."""
    with open(path, "rb") as f:
        head = f.read(8)

    if head.startswith(codecs.BOM_UTF16_LE) or head.startswith(codecs.BOM_UTF16_BE):
        encodings = ("utf-16", "utf-8-sig", "cp1252", "latin-1")
    elif head.startswith(codecs.BOM_UTF8):
        encodings = ("utf-8-sig", "utf-8", "cp1252", "latin-1")
    elif b"\x00" in head:
        # UTF-16 sin BOM: los textos ASCII tienen NUL en posiciones
        # pares o impares según el endianness.
        even_nuls = head[0::2].count(0)
        odd_nuls  = head[1::2].count(0)
        encodings = (
            "utf-16-le" if odd_nuls >= even_nuls else "utf-16-be",
            "utf-16-le", "utf-16-be",
            "utf-8", "cp1252", "latin-1",
        )
    else:
        encodings = ("utf-8", "utf-8-sig", "cp1252", "latin-1")

    for enc in encodings:
        try:
            with open(path, "r", encoding=enc) as f:
                return f.read()
        except (UnicodeDecodeError, UnicodeError):
            continue

    with open(path, "r", encoding="utf-8", errors="replace") as f:
        return f.read()


def _split_by_marker(paragraphs: list) -> list:
    """Divide una lista de párrafos en capítulos usando
    CHAPTER_MARKER_REGEX. Sin marcadores → 1 solo capítulo."""
    chapters, current, found = [], [], False
    for p in paragraphs:
        if C.CHAPTER_MARKER_REGEX.match(p.strip()):
            found = True
            if current:
                chapters.append(current)
            current = [p]
        else:
            current.append(p)
    if current:
        chapters.append(current)
    return chapters if found else [paragraphs]


def load_txt(path: str) -> list:
    raw = _read_txt_raw(path)
    paragraphs = [p.strip() for p in re.split(r"\n\s*\n", raw) if p.strip()]
    chapters = _split_by_marker(paragraphs)
    if len(chapters) > 1:
        print(f"✅ Capítulos detectados en TXT por marcador de texto: {len(chapters)}")
    else:
        print("ℹ  Sin marcadores de capítulo en TXT — se trata como 1 solo capítulo.")
    return chapters


# ==========================================================
# DOCX
# ==========================================================

def _is_title_style(style) -> bool:
    """True para «Title»/«Subtitle»: metadatos del libro, no capítulos."""
    if style is None:
        return False
    sid  = (getattr(style, "style_id", "") or "").strip().lower()
    name = (getattr(style, "name", "") or "").strip().lower()
    return (sid in ("title", "subtitle")
            or name in ("title", "subtitle", "título", "subtítulo",
                        "titulo", "subtitulo"))


def _docx_paragraphs(path: str) -> list:
    return [
        p.text.strip()
        for p in Document(path).paragraphs
        if p.text.strip() and not _is_title_style(p.style)
    ]


def _heading_level(style) -> int | None:
    """Devuelve 1 o 2 para estilos de encabezado de capítulo.

    Se usa `style_id` (que Word mantiene como Heading1/Heading2
    incluso en documentos en español) y, como respaldo, el nombre
    traducido. «Title» se excluye a propósito: suele ser el título
    del libro, no el de cada capítulo.
    """
    if style is None:
        return None
    sid  = (getattr(style, "style_id", "") or "").strip().lower()
    name = (getattr(style, "name", "") or "").strip().lower()

    level1_names = {
        "heading 1", "heading1", "título 1", "titulo 1",
        "encabezado 1", "encabezado1", "titre 1", "überschrift 1",
    }
    level2_names = {
        "heading 2", "heading2", "título 2", "titulo 2",
        "encabezado 2", "encabezado2", "titre 2", "überschrift 2",
    }

    if sid in ("heading1", "heading 1") or name in level1_names:
        return 1
    if sid in ("heading2", "heading 2") or name in level2_names:
        return 2
    return None


def _docx_by_heading(path: str) -> list:
    """Divide un DOCX por estilos de encabezado de Word.

    Si hay varios niveles, usa el que tenga más encabezados (es el
    que probablemente corresponde a los capítulos). Así un libro con
    «Title» + «Heading 1» no se parte en dos mitades gigantes, y un
    libro con partes en Heading 1 y capítulos en Heading 2 se divide
    por capítulos.
    """
    doc = Document(path)
    counts: dict[int, int] = {}
    for p in doc.paragraphs:
        if not p.text.strip():
            continue
        level = _heading_level(p.style)
        if level:
            counts[level] = counts.get(level, 0) + 1

    if not counts:
        return []

    candidates = [(lvl, cnt) for lvl, cnt in counts.items() if cnt >= 2]
    if candidates:
        # Más encabezados primero; en empate, el nivel más alto (1).
        heading_level = max(candidates, key=lambda x: (x[1], -x[0]))[0]
    else:
        # Un solo encabezado: no hay forma fiable de dividir.
        return []

    chapters, current, current_styles = [], [], []
    for p in doc.paragraphs:
        text = p.text.strip()
        if not text or _is_title_style(p.style):
            continue

        if _heading_level(p.style) == heading_level:
            if current:
                # No convertir en capítulo un prólogo que solo contiene
                # el título del libro (Title) o una cabecera de otro nivel.
                only_heading = (len(current) == 1
                                and _heading_level(current_styles[0]) is not None)
                if not only_heading:
                    chapters.append(current)
            current, current_styles = [text], [p.style]
        else:
            current.append(text)
            current_styles.append(p.style)

    if current:
        only_heading = (len(current) == 1
                        and _heading_level(current_styles[0]) is not None)
        if not only_heading:
            chapters.append(current)
    return chapters


def load_docx(path: str) -> list:
    chapters = _docx_by_heading(path)
    if len(chapters) > 1:
        print(f"✅ Capítulos detectados en DOCX por estilo de encabezado: {len(chapters)}")
        return chapters
    paragraphs = _docx_paragraphs(path)
    chapters = _split_by_marker(paragraphs)
    if len(chapters) > 1:
        print(f"✅ Capítulos detectados en DOCX por marcador de texto: {len(chapters)}")
    else:
        print("ℹ  Sin capítulos detectados en DOCX — se trata como 1 solo capítulo.\n"
              "   -N/--chapters-per-file no tendrá efecto.")
    return chapters


# ==========================================================
# EPUB — lector nativo
# ==========================================================

_BLOCK_TAGS = ["p", "h1", "h2", "h3", "h4", "h5", "h6",
               "li", "blockquote", "div", "td", "dt", "dd", "pre"]


def _epub_paragraphs(soup: BeautifulSoup) -> list:
    """Extrae párrafos sin duplicar texto de divs que contienen <p>.

    Se procesan solo los bloques "hoja": si un div contiene párrafos,
    se usan los párrafos; si contiene texto suelto, se usa el div.
    """
    paragraphs = []
    for el in soup.find_all(_BLOCK_TAGS):
        if el.find(_BLOCK_TAGS) is not None:
            continue
        text = el.get_text(" ", strip=True)
        if text:
            paragraphs.append(text)

    if not paragraphs:
        text = soup.get_text(separator="\n", strip=True)
        paragraphs = [p.strip() for p in text.splitlines() if p.strip()]
    return paragraphs


def _epub_native(path: str) -> list:
    """Lee el EPUB usando book.spine (orden real de lectura),
    ignorando items linear='no' y evitando duplicados."""
    book = epub.read_epub(path)
    items_by_id = {
        item.get_id(): item
        for item in book.get_items()
        if isinstance(item, epub.EpubHtml) and not isinstance(item, epub.EpubNav)
    }
    chapters, seen = [], set()
    for spine_id, linear in book.spine:
        if linear in ("no", False) or spine_id in seen:
            continue
        seen.add(spine_id)
        item = items_by_id.get(spine_id)
        if not item:
            continue
        soup = BeautifulSoup(item.get_content(), "html.parser")
        paragraphs = _epub_paragraphs(soup)
        if paragraphs:
            chapters.append(paragraphs)
    return chapters


# ==========================================================
# EPUB — vía Calibre
# ==========================================================

def _find_calibre() -> str | None:
    for path in C.CALIBRE_EBOOK_CONVERT_CANDIDATES:
        if os.path.exists(path):
            return path
    return shutil.which("ebook-convert")


def _calibre_convert(epub_path: str, out_docx: str) -> bool:
    exe = _find_calibre()
    if not exe:
        return False
    try:
        subprocess.run(
            [exe, epub_path, out_docx],
            check=True,
            stdout=subprocess.DEVNULL,
            stderr=subprocess.DEVNULL,
        )
        return os.path.exists(out_docx) and os.path.getsize(out_docx) > 0
    except Exception:
        return False


def load_epub(path: str) -> list:
    """Intenta primero vía Calibre (mejor detección de capítulos),
    cae al lector nativo si Calibre no está disponible o no detecta
    más de un capítulo por estilos de encabezado."""
    if C.USE_CALIBRE_FOR_EPUB and _find_calibre():
        tmp = tempfile.mkdtemp(prefix="superaudio_calibre_")
        try:
            print("📐 Convirtiendo EPUB con Calibre …")
            out_docx = os.path.join(tmp, "conv.docx")
            if _calibre_convert(path, out_docx):
                chapters = _docx_by_heading(out_docx)
                if len(chapters) > 1:
                    print(f"✅ Capítulos via Calibre: {len(chapters)}")
                    return chapters
                print("ℹ  Calibre no detectó capítulos por encabezado → lector nativo.")
            else:
                print("ℹ  Falló la conversión con Calibre → lector nativo.")
        except Exception as e:
            print(f"⚠  Error Calibre ({e}) → lector nativo.")
        finally:
            shutil.rmtree(tmp, ignore_errors=True)
    else:
        if C.USE_CALIBRE_FOR_EPUB:
            print("ℹ  Calibre no encontrado → lector nativo.")

    return _epub_native(path)


# ==========================================================
# TOC del EPUB
# ==========================================================

def _flatten_toc(items) -> list:
    """Aplana book.toc, que puede contener Link, Section y tuplas
    (Section, hijos)."""
    flat = []
    for it in items or []:
        if isinstance(it, tuple):
            # Formato (Section, [children])
            if len(it) == 2 and isinstance(it[1], (list, tuple)):
                flat.append(it[0])
                flat.extend(_flatten_toc(it[1]))
            else:
                flat.extend(_flatten_toc(it))
            continue

        flat.append(it)
        children = getattr(it, "children", None)
        if children:
            flat.extend(_flatten_toc(children))
    return flat


def _toc_titles(path: str, expected: int) -> list[str]:
    """Devuelve títulos del TOC solo si el número coincide.

    Si el TOC tiene más/menos entradas que capítulos detectados, es
    preferible usar «Capítulo N» antes que asignar títulos corridos.
    """
    try:
        book = epub.read_epub(path)
        titles = []
        for it in _flatten_toc(book.toc):
            title = getattr(it, "title", None)
            if title and str(title).strip():
                titles.append(str(title).strip())
        if len(titles) == expected and titles:
            return titles
    except Exception:
        pass
    return []


# ==========================================================
# Función pública principal
# ==========================================================

def load_book(path: str) -> tuple[list, list]:
    """Carga el libro y devuelve (chapters_clean, chapter_titles).

    chapters_clean : lista de capítulos; cada capítulo es una
                     lista de párrafos ya limpios.
    chapter_titles : lista de títulos (strings) del mismo largo.
    """
    ext = Path(path).suffix.lower()

    if ext == ".txt":
        raw_chapters = load_txt(path)
    elif ext == ".docx":
        raw_chapters = load_docx(path)
    elif ext == ".epub":
        raw_chapters = load_epub(path)
    else:
        raise ValueError(f"Formato no soportado: {ext}")

    # Limpiar párrafos
    chapters_clean = []
    for paragraphs in raw_chapters:
        cleaned = [clean(p) for p in paragraphs]
        cleaned = [p for p in cleaned if p]
        if cleaned:
            chapters_clean.append(cleaned)

    # Títulos por defecto
    titles = [f"Capítulo {i + 1}" for i in range(len(chapters_clean))]

    # Títulos reales del TOC del EPUB (solo si cuadran exactamente)
    if ext == ".epub":
        real = _toc_titles(path, len(titles))
        if real:
            titles = real
            print("✅ Títulos del TOC del EPUB aplicados.")

    return chapters_clean, titles
